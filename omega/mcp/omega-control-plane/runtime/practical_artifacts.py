#!/usr/bin/env python3
import argparse, json, os, sys, math
from pathlib import Path

def result(**kw):
    print(json.dumps(kw, ensure_ascii=False))

def doctor():
    mods={}
    for name in ['docx','reportlab','xlsxwriter','pptx','h5py']:
        try:
            m=__import__(name); mods[name]={'available':True,'version':getattr(m,'__version__',None)}
        except Exception as e: mods[name]={'available':False,'error':str(e)}
    return {'operation':'doctor','modules':mods,'ready':all(x['available'] for x in mods.values())}

def ensure_parent(path): Path(path).parent.mkdir(parents=True, exist_ok=True)

def make_docx(req):
    from docx import Document
    from docx.shared import Pt, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    out=req['output']; ensure_parent(out)
    d=Document(); sec=d.sections[0]
    margins=req.get('marginsCm',{})
    sec.top_margin=Cm(float(margins.get('top',2.0)));sec.bottom_margin=Cm(float(margins.get('bottom',2.0)))
    sec.left_margin=Cm(float(margins.get('left',2.0)));sec.right_margin=Cm(float(margins.get('right',2.0)))
    font=req.get('font',{}); style=d.styles['Normal'];style.font.name=font.get('name','Arial');style.font.size=Pt(float(font.get('size',11)))
    if req.get('title'):
        p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;r=p.add_run(str(req['title']));r.bold=True;r.font.size=Pt(float(font.get('titleSize',18)))
    for block in req.get('blocks',[]):
        if block.get('type')=='heading': d.add_heading(str(block.get('text','')), level=max(1,min(9,int(block.get('level',1)))))
        elif block.get('type')=='table':
            rows=block.get('rows',[]); cols=max([len(r) for r in rows], default=1); t=d.add_table(rows=0, cols=cols); t.style='Table Grid'
            for row in rows:
                cells=t.add_row().cells
                for i,v in enumerate(row): cells[i].text=str(v)
        else: d.add_paragraph(str(block.get('text','')))
    if req.get('footer'):
        sec.footer.paragraphs[0].text=str(req['footer'])
    d.save(out)
    return {'kind':'docx','output':out,'formattingPreserved':True}

def make_pdf(req):
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import cm
    out=req['output'];ensure_parent(out); styles=getSampleStyleSheet(); margins=req.get('marginsCm',{})
    doc=SimpleDocTemplate(out,pagesize=A4,leftMargin=float(margins.get('left',2))*cm,rightMargin=float(margins.get('right',2))*cm,topMargin=float(margins.get('top',2))*cm,bottomMargin=float(margins.get('bottom',2))*cm)
    story=[]
    if req.get('title'): story += [Paragraph(str(req['title']),styles['Title']),Spacer(1,8)]
    for b in req.get('blocks',[]):
        if b.get('type')=='heading': story += [Paragraph(str(b.get('text','')),styles['Heading2']),Spacer(1,5)]
        elif b.get('type')=='table':
            t=Table([[str(x) for x in r] for r in b.get('rows',[])], repeatRows=1);t.setStyle(TableStyle([('GRID',(0,0),(-1,-1),0.25,colors.grey),('VALIGN',(0,0),(-1,-1),'TOP')]));story.append(t)
        else: story += [Paragraph(str(b.get('text','')),styles['BodyText']),Spacer(1,5)]
    footer=str(req.get('footer',''))
    def on_page(canvas,doc):
        if footer: canvas.saveState();canvas.setFont('Helvetica',8);canvas.drawString(doc.leftMargin,1*cm,footer);canvas.restoreState()
    doc.build(story,onFirstPage=on_page,onLaterPages=on_page)
    return {'kind':'pdf','output':out,'formattingPreserved':True}


def make_news_pdf(req):
    from reportlab.platypus import BaseDocTemplate, Frame, PageTemplate, Paragraph, Spacer, KeepTogether
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.units import cm
    from reportlab.lib import colors
    out=req['output']; ensure_parent(out)
    page_w,page_h=A4; left=1.4*cm; right=1.4*cm; top=2.6*cm; bottom=1.5*cm; gap=.55*cm
    col=(page_w-left-right-gap)/2
    doc=BaseDocTemplate(out,pagesize=A4,leftMargin=left,rightMargin=right,topMargin=top,bottomMargin=bottom)
    f1=Frame(left,bottom,col,page_h-top-bottom,id='col1',showBoundary=0)
    f2=Frame(left+col+gap,bottom,col,page_h-top-bottom,id='col2',showBoundary=0)
    masthead=str(req.get('brand','Mojealterego News')); issue=str(req.get('issue','')); period=str(req.get('period',''))
    def header(canvas,doc):
        canvas.saveState(); canvas.setFillColor(colors.HexColor('#111111')); canvas.rect(0,page_h-2.15*cm,page_w,2.15*cm,fill=1,stroke=0)
        canvas.setFillColor(colors.white); canvas.setFont('Helvetica-Bold',22); canvas.drawCentredString(page_w/2,page_h-1.05*cm,masthead)
        canvas.setFont('Helvetica',8); canvas.drawCentredString(page_w/2,page_h-1.62*cm,'  •  '.join([x for x in [issue,period] if x]))
        canvas.setFillColor(colors.HexColor('#444444')); canvas.setFont('Helvetica',7); canvas.drawString(left,.75*cm,str(req.get('footer','Mojealterego News')))
        canvas.drawRightString(page_w-right,.75*cm,f'Strona {doc.page}'); canvas.restoreState()
    doc.addPageTemplates([PageTemplate(id='news',frames=[f1,f2],onPage=header)])
    styles=getSampleStyleSheet()
    title=ParagraphStyle('NewsTitle',parent=styles['Title'],fontName='Helvetica-Bold',fontSize=17,leading=19,spaceAfter=8,textColor=colors.HexColor('#111111'))
    dek=ParagraphStyle('NewsDek',parent=styles['BodyText'],fontName='Helvetica-Oblique',fontSize=9.5,leading=12,spaceAfter=10,textColor=colors.HexColor('#444444'))
    section=ParagraphStyle('NewsSection',parent=styles['Heading2'],fontName='Helvetica-Bold',fontSize=11,leading=13,spaceBefore=7,spaceAfter=4,textColor=colors.HexColor('#111111'),borderColor=colors.HexColor('#888888'),borderWidth=.4,borderPadding=(0,0,2,0))
    head=ParagraphStyle('NewsHead',parent=styles['Heading3'],fontName='Helvetica-Bold',fontSize=9.5,leading=11,spaceAfter=2)
    body=ParagraphStyle('NewsBody',parent=styles['BodyText'],fontName='Helvetica',fontSize=8.3,leading=10.5,spaceAfter=6,alignment=0)
    story=[]
    if req.get('title'): story.append(Paragraph(str(req['title']),title))
    if req.get('dek'): story.append(Paragraph(str(req['dek']),dek))
    for sec in req.get('sections',[]):
        story.append(Paragraph(str(sec.get('title','')),section))
        for item in sec.get('items',[]):
            story.append(KeepTogether([Paragraph(str(item.get('title','')),head),Paragraph(str(item.get('summary','')),body)]))
    doc.build(story)
    return {'kind':'pdf','output':out,'formattingPreserved':True,'layout':'newspaper-two-column'}

def make_xlsx(req):
    import xlsxwriter
    out=req['output'];ensure_parent(out); wb=xlsxwriter.Workbook(out)
    header=wb.add_format({'bold':True,'border':1,'bg_color':'#E7E6E6'}); cell=wb.add_format({'border':1});
    tables=[]
    for si,s in enumerate(req.get('sheets',[]) or [{'name':'Sheet1','rows':[]}]):
        ws=wb.add_worksheet(str(s.get('name',f'Sheet{si+1}'))[:31]); rows=s.get('rows',[])
        for r,row in enumerate(rows):
            for c,val in enumerate(row):
                fmt=header if r==0 and s.get('header',True) else cell
                if isinstance(val,dict) and 'formula' in val: ws.write_formula(r,c,str(val['formula']),fmt,val.get('value',0))
                else: ws.write(r,c,val,fmt)
        if s.get('table') and rows:
            name=s.get('tableName',f'Table{si+1}').replace(' ','_'); ws.add_table(0,0,len(rows)-1,max(0,len(rows[0])-1),{'name':name,'columns':[{'header':str(x)} for x in rows[0]]});tables.append(name)
        widths=s.get('widths',[])
        for c,w in enumerate(widths): ws.set_column(c,c,float(w))
    wb.close()
    return {'kind':'xlsx','output':out,'formulasWritten':True,'tables':tables,'pivotTablesCreated':False,'pivotNote':'Native pivot-table creation is not claimed by this runtime.'}

def make_pptx(req):
    from pptx import Presentation
    from pptx.util import Inches, Pt
    out=req['output'];ensure_parent(out); prs=Presentation()
    for s in req.get('slides',[]):
        layout=prs.slide_layouts[1]; slide=prs.slides.add_slide(layout); slide.shapes.title.text=str(s.get('title',''))
        tf=slide.placeholders[1].text_frame; tf.clear();
        for i,line in enumerate(s.get('bullets',[])):
            p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=str(line);p.font.size=Pt(float(s.get('fontSize',20)))
    if not req.get('slides'): prs.slides.add_slide(prs.slide_layouts[0])
    prs.save(out); return {'kind':'pptx','output':out,'slides':len(prs.slides)}

def h5_summary(req):
    import h5py
    path=req['path']; max_items=int(req.get('maxItems',1000)); items=[]
    with h5py.File(path,'r') as h:
        def visit(name,obj):
            if len(items)>=max_items:return
            row={'path':'/'+name,'type':'dataset' if isinstance(obj,h5py.Dataset) else 'group'}
            if isinstance(obj,h5py.Dataset): row.update({'shape':list(obj.shape),'dtype':str(obj.dtype),'size':int(obj.size)})
            items.append(row)
        h.visititems(visit)
    return {'kind':'hdf5','path':path,'items':items,'truncated':len(items)>=max_items}

def sequence_summary(req):
    path=Path(req['path']); max_records=int(req.get('maxRecords',100000)); n=0;bases=0;gc=0;lengths=[]
    with path.open('rt',errors='replace') as f:
        first=f.readline();f.seek(0)
        if first.startswith('>'):
            seq=[]
            for line in f:
                if line.startswith('>'):
                    if seq:
                        s=''.join(seq).strip().upper();n+=1;bases+=len(s);gc+=s.count('G')+s.count('C');lengths.append(len(s));seq=[]
                        if n>=max_records:break
                else: seq.append(line.strip())
            if seq and n<max_records:
                s=''.join(seq).strip().upper();n+=1;bases+=len(s);gc+=s.count('G')+s.count('C');lengths.append(len(s))
            fmt='FASTA'
        elif first.startswith('@'):
            while n<max_records:
                h=f.readline();
                if not h:break
                seq=f.readline().strip().upper();plus=f.readline();qual=f.readline()
                if not qual:break
                n+=1;bases+=len(seq);gc+=seq.count('G')+seq.count('C');lengths.append(len(seq))
            fmt='FASTQ'
        else: raise ValueError('Unknown sequencing text format')
    return {'kind':'sequence','format':fmt,'records':n,'bases':bases,'gcFraction':(gc/bases if bases else 0),'meanLength':(sum(lengths)/len(lengths) if lengths else 0),'truncated':n>=max_records}

def main():
    ap=argparse.ArgumentParser();ap.add_argument('--request',required=True);args=ap.parse_args();req=json.loads(Path(args.request).read_text('utf8'));op=req.get('operation')
    try:
        if op=='doctor': out=doctor()
        elif op=='docx': out=make_docx(req)
        elif op=='pdf': out=make_pdf(req)
        elif op=='news-pdf': out=make_news_pdf(req)
        elif op=='xlsx': out=make_xlsx(req)
        elif op=='pptx': out=make_pptx(req)
        elif op=='h5-summary': out=h5_summary(req)
        elif op=='sequence-summary': out=sequence_summary(req)
        else: raise ValueError(f'Unsupported operation: {op}')
        result(ok=True,**out)
    except Exception as e:
        result(ok=False,error=f'{type(e).__name__}: {e}');sys.exit(1)
if __name__=='__main__':main()
